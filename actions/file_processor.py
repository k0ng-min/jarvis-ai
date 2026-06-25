# actions/file_processor.py — 자비스 파일 처리기 (Claude CLI 기반)

from __future__ import annotations

import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

_OS = platform.system()

# ─── Claude CLI 헬퍼 ─────────────────────────────────────────────────────────

def _call_claude(prompt: str, timeout: int = 60) -> str:
    try:
        result = subprocess.run(
            ["claude", "-p", prompt, "--output-format", "json"],
            capture_output=True, text=True, timeout=timeout, encoding="utf-8"
        )
        if result.returncode == 0:
            data = json.loads(result.stdout.strip())
            return data.get("result", "").strip()
    except Exception as e:
        print(f"[파일처리] Claude 오류: {e}")
    return ""


def _call_claude_with_file(prompt: str, file_path: str, timeout: int = 90) -> str:
    """파일 경로를 포함한 Claude 호출 (이미지/바이너리용)"""
    full_prompt = f"파일: {file_path}\n\n{prompt}"
    return _call_claude(full_prompt, timeout)


# ─── 의존성 설치 헬퍼 ────────────────────────────────────────────────────────

def _ensure(*packages: str) -> bool:
    for pkg in packages:
        mod = pkg.split("[")[0].replace("-", "_")
        try:
            __import__(mod)
        except ImportError:
            subprocess.run(
                [sys.executable, "-m", "pip", "install", pkg, "--break-system-packages", "-q"],
                capture_output=True
            )
    return True


# ─── 카테고리 감지 ───────────────────────────────────────────────────────────

_EXT_CAT = {
    **dict.fromkeys(["jpg","jpeg","png","gif","webp","bmp","tiff","svg","ico"], "image"),
    **dict.fromkeys(["mp4","avi","mov","mkv","wmv","flv","webm","m4v"],         "video"),
    **dict.fromkeys(["mp3","wav","ogg","m4a","aac","flac","wma","opus"],        "audio"),
    **dict.fromkeys(["pdf"],                                                     "pdf"),
    **dict.fromkeys(["doc","docx"],                                              "word"),
    **dict.fromkeys(["xls","xlsx","ods"],                                        "excel"),
    **dict.fromkeys(["ppt","pptx"],                                              "pptx"),
    **dict.fromkeys(["py","js","ts","jsx","tsx","html","css","java","c","cpp",
                     "cs","go","rs","rb","php","swift","kt","sh","sql","lua"],   "code"),
    **dict.fromkeys(["zip","rar","tar","gz","7z","bz2","xz"],                   "archive"),
    **dict.fromkeys(["txt","md","rst","log"],                                    "text"),
    **dict.fromkeys(["csv","tsv"],                                               "csv"),
    **dict.fromkeys(["json"],                                                    "json"),
    **dict.fromkeys(["xml"],                                                     "xml"),
}

def _cat(path: Path) -> str:
    return _EXT_CAT.get(path.suffix.lower().lstrip("."), "unknown")


# ─── 이미지 처리 ─────────────────────────────────────────────────────────────

def _image_describe(path: Path) -> str:
    return _call_claude_with_file(
        f"이 이미지를 자세히 설명하세요. 내용, 색상, 구성 요소, 중요한 세부 사항을 포함하세요.",
        str(path)
    ) or "이미지를 설명할 수 없습니다."


def _image_ocr(path: Path) -> str:
    result = _call_claude_with_file(
        f"이 이미지에서 모든 텍스트를 추출하세요. 원본 형식을 최대한 유지하세요.",
        str(path)
    )
    return result or "텍스트를 추출할 수 없습니다."


def _image_resize(path: Path, params: dict) -> str:
    _ensure("Pillow")
    from PIL import Image
    width  = int(params.get("width",  0))
    height = int(params.get("height", 0))
    if not width and not height:
        return "너비 또는 높이를 지정해주세요."
    img = Image.open(path)
    ow, oh = img.size
    if width and not height:
        height = int(oh * width / ow)
    elif height and not width:
        width = int(ow * height / oh)
    img = img.resize((width, height), Image.LANCZOS)
    out = path.parent / f"{path.stem}_resized{path.suffix}"
    img.save(out)
    return f"이미지 크기 조정 완료: {width}×{height} → {out.name}"


def _image_convert(path: Path, params: dict) -> str:
    _ensure("Pillow")
    from PIL import Image
    fmt = params.get("format", "PNG").upper()
    ext_map = {"JPG": "JPEG", "JPEG": "JPEG", "PNG": "PNG", "WEBP": "WEBP",
               "BMP": "BMP", "TIFF": "TIFF", "ICO": "ICO"}
    pil_fmt = ext_map.get(fmt, fmt)
    img = Image.open(path).convert("RGBA" if pil_fmt == "PNG" else "RGB")
    ext = fmt.lower().replace("jpeg", "jpg")
    out = path.parent / f"{path.stem}.{ext}"
    img.save(out, format=pil_fmt)
    return f"이미지 변환 완료: {out.name}"


def _image_compress(path: Path, params: dict) -> str:
    _ensure("Pillow")
    from PIL import Image
    quality = int(params.get("quality", 75))
    img = Image.open(path)
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    out = path.parent / f"{path.stem}_compressed.jpg"
    img.save(out, "JPEG", quality=quality, optimize=True)
    orig  = path.stat().st_size
    comp  = out.stat().st_size
    saved = (1 - comp / orig) * 100 if orig else 0
    return f"이미지 압축 완료: {out.name} ({saved:.1f}% 절약)"


def _process_image(path: Path, action: str, params: dict) -> str:
    actions = {
        "describe": lambda: _image_describe(path),
        "ocr":      lambda: _image_ocr(path),
        "resize":   lambda: _image_resize(path, params),
        "convert":  lambda: _image_convert(path, params),
        "compress": lambda: _image_compress(path, params),
    }
    fn = actions.get(action)
    if fn:
        return fn()
    return _image_describe(path)


# ─── PDF 처리 ─────────────────────────────────────────────────────────────────

def _pdf_extract_text(path: Path) -> str:
    _ensure("pymupdf")
    try:
        import fitz  # pymupdf
        doc = fitz.open(str(path))
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text()
            if text.strip():
                pages.append(f"[페이지 {i+1}]\n{text}")
        return "\n\n".join(pages[:30])  # 최대 30페이지
    except Exception as e:
        return f"PDF 텍스트 추출 실패: {e}"


def _pdf_summarize(path: Path) -> str:
    text = _pdf_extract_text(path)
    if not text or text.startswith("PDF"):
        return "PDF 텍스트를 추출할 수 없습니다."
    return _call_claude(
        f"다음 PDF의 핵심 내용을 한국어로 요약하세요:\n\n{text[:8000]}"
    ) or "요약할 수 없습니다."


def _pdf_to_word(path: Path) -> str:
    _ensure("pymupdf", "python-docx")
    import fitz
    from docx import Document
    doc_in  = fitz.open(str(path))
    doc_out = Document()
    for i, page in enumerate(doc_in):
        if i > 0:
            doc_out.add_page_break()
        text = page.get_text()
        for para in text.split("\n"):
            if para.strip():
                doc_out.add_paragraph(para)
    out = path.parent / f"{path.stem}.docx"
    doc_out.save(str(out))
    return f"PDF를 Word로 변환 완료: {out.name}"


def _process_pdf(path: Path, action: str, params: dict) -> str:
    actions = {
        "summarize":    lambda: _pdf_summarize(path),
        "extract_text": lambda: _pdf_extract_text(path),
        "to_word":      lambda: _pdf_to_word(path),
    }
    fn = actions.get(action)
    return fn() if fn else _pdf_summarize(path)


# ─── Word/텍스트/마크다운 처리 ───────────────────────────────────────────────

def _read_text_file(path: Path) -> str:
    try:
        if path.suffix.lower() in (".docx", ".doc"):
            _ensure("python-docx")
            from docx import Document
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"파일 읽기 실패: {e}"


def _process_text(path: Path, action: str, params: dict) -> str:
    content = _read_text_file(path)
    if not content:
        return "파일 내용을 읽을 수 없습니다."

    if action == "summarize":
        return _call_claude(f"다음을 한국어로 요약하세요:\n\n{content[:8000]}") or "요약 실패"
    elif action == "reformat":
        fmt = params.get("format", "markdown")
        return _call_claude(
            f"다음 텍스트를 {fmt} 형식으로 재구성하세요:\n\n{content[:6000]}"
        ) or "재구성 실패"
    else:
        return _call_claude(f"다음 파일을 분석하세요:\n\n{content[:8000]}") or "분석 실패"


# ─── CSV/Excel 처리 ──────────────────────────────────────────────────────────

def _read_csv(path: Path) -> str:
    """CSV/Excel 내용을 텍스트로"""
    try:
        if path.suffix.lower() in (".xlsx", ".xls", ".ods"):
            _ensure("openpyxl")
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            rows = []
            for sheet in wb.sheetnames[:3]:
                ws = wb[sheet]
                rows.append(f"[시트: {sheet}]")
                for row in list(ws.iter_rows(values_only=True))[:100]:
                    rows.append(",".join(str(v or "") for v in row))
            return "\n".join(rows)
        else:
            return path.read_text(encoding="utf-8-sig", errors="replace")
    except Exception as e:
        return f"파일 읽기 실패: {e}"


def _process_spreadsheet(path: Path, action: str, params: dict) -> str:
    content = _read_csv(path)
    if not content:
        return "스프레드시트를 읽을 수 없습니다."

    if action == "analyze":
        return _call_claude(
            f"이 데이터를 분석하고 주요 인사이트를 한국어로 제공하세요:\n\n{content[:6000]}"
        ) or "분석 실패"
    elif action == "filter":
        condition = params.get("condition", "")
        return _call_claude(
            f"조건 '{condition}'에 맞는 행을 필터링하세요:\n\n{content[:6000]}"
        ) or "필터링 실패"
    elif action == "sort":
        col = params.get("column", "")
        return _call_claude(
            f"'{col}' 기준으로 정렬하세요:\n\n{content[:6000]}"
        ) or "정렬 실패"
    elif action == "convert":
        fmt = params.get("format", "json")
        result = _call_claude(
            f"이 CSV 데이터를 {fmt} 형식으로 변환하세요:\n\n{content[:6000]}"
        )
        if result:
            out = path.parent / f"{path.stem}.{fmt}"
            out.write_text(result, encoding="utf-8")
            return f"변환 완료: {out.name}"
        return "변환 실패"
    else:
        return _call_claude(f"이 데이터를 분석하세요:\n\n{content[:6000]}") or "분석 실패"


# ─── JSON 처리 ───────────────────────────────────────────────────────────────

def _process_json(path: Path, action: str, params: dict) -> str:
    content = path.read_text(encoding="utf-8", errors="replace")

    if action == "validate":
        try:
            data = json.loads(content)
            return f"✅ 유효한 JSON입니다. 키 수: {len(data) if isinstance(data, dict) else len(data) if isinstance(data, list) else 1}"
        except json.JSONDecodeError as e:
            return f"❌ 유효하지 않은 JSON: {e}"
    elif action == "format":
        try:
            data = json.loads(content)
            out = path.parent / f"{path.stem}_formatted.json"
            out.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
            return f"JSON 포맷 완료: {out.name}"
        except Exception as e:
            return f"포맷 실패: {e}"
    elif action == "extract":
        key = params.get("key", "")
        return _call_claude(
            f"이 JSON에서 '{key}' 관련 데이터를 추출하세요:\n\n{content[:6000]}"
        ) or "추출 실패"
    else:
        return _call_claude(
            f"이 JSON 파일을 분석하고 설명하세요:\n\n{content[:6000]}"
        ) or "분석 실패"


# ─── 코드 처리 ───────────────────────────────────────────────────────────────

def _process_code(path: Path, action: str, params: dict) -> str:
    content = path.read_text(encoding="utf-8", errors="replace")
    lang    = path.suffix.lstrip(".")

    if action == "explain":
        return _call_claude(
            f"이 {lang} 코드를 한국어로 설명하세요:\n\n```{lang}\n{content[:6000]}\n```"
        ) or "설명 실패"
    elif action == "review":
        return _call_claude(
            f"이 {lang} 코드를 리뷰하고 개선 사항을 한국어로 제안하세요:\n\n```{lang}\n{content[:6000]}\n```"
        ) or "리뷰 실패"
    elif action == "fix":
        issue = params.get("issue", "버그 수정 및 개선")
        result = _call_claude(
            f"이 {lang} 코드의 문제를 수정하세요 ({issue}):\n\n```{lang}\n{content[:6000]}\n```\n수정된 코드만 반환하세요."
        )
        if result:
            # 코드 블록 제거
            result = re.sub(r"```\w*\n?", "", result).rstrip("`").strip()
            out = path.parent / f"{path.stem}_fixed{path.suffix}"
            out.write_text(result, encoding="utf-8")
            return f"코드 수정 완료: {out.name}"
        return "수정 실패"
    elif action == "run":
        if lang == "py":
            try:
                r = subprocess.run(
                    [sys.executable, str(path)],
                    capture_output=True, text=True, timeout=30
                )
                out = r.stdout.strip() or r.stderr.strip()
                return f"실행 결과:\n{out[:2000]}" if out else "출력 없음"
            except subprocess.TimeoutExpired:
                return "실행 시간 초과 (30초)"
            except Exception as e:
                return f"실행 실패: {e}"
        return f".{lang} 파일 실행은 지원되지 않습니다."
    elif action == "document":
        result = _call_claude(
            f"이 {lang} 코드에 한국어 문서 주석을 추가하세요:\n\n```{lang}\n{content[:6000]}\n```\n주석이 추가된 코드만 반환하세요."
        )
        if result:
            result = re.sub(r"```\w*\n?", "", result).rstrip("`").strip()
            out = path.parent / f"{path.stem}_documented{path.suffix}"
            out.write_text(result, encoding="utf-8")
            return f"문서화 완료: {out.name}"
        return "문서화 실패"
    else:
        return _call_claude(
            f"이 {lang} 코드를 분석하세요:\n\n```{lang}\n{content[:6000]}\n```"
        ) or "분석 실패"


# ─── 오디오 처리 ─────────────────────────────────────────────────────────────

def _process_audio(path: Path, action: str, params: dict) -> str:
    if action == "transcribe":
        # speech_recognition으로 WAV 전사, 다른 포맷은 ffmpeg으로 변환 후
        _ensure("SpeechRecognition")
        import speech_recognition as sr

        wav_path = path
        if path.suffix.lower() != ".wav":
            wav_path = path.parent / f"{path.stem}_tmp.wav"
            try:
                subprocess.run(
                    ["ffmpeg", "-i", str(path), "-ac", "1", "-ar", "16000",
                     str(wav_path), "-y", "-loglevel", "quiet"],
                    timeout=60, check=True
                )
            except Exception as e:
                return f"오디오 변환 실패 (ffmpeg 필요): {e}"

        rec = sr.Recognizer()
        try:
            with sr.AudioFile(str(wav_path)) as src:
                audio = rec.record(src)
            text = rec.recognize_google(audio, language="ko-KR")
            return f"전사 결과:\n{text}"
        except sr.UnknownValueError:
            return "오디오를 인식할 수 없습니다."
        except Exception as e:
            return f"전사 실패: {e}"
        finally:
            if wav_path != path and wav_path.exists():
                wav_path.unlink(missing_ok=True)

    elif action == "trim":
        start = params.get("start", "0")
        end   = params.get("end", "")
        out   = path.parent / f"{path.stem}_trimmed{path.suffix}"
        end_args = ["-to", end] if end else []
        try:
            subprocess.run(
                ["ffmpeg", "-i", str(path), "-ss", start] + end_args +
                ["-c", "copy", str(out), "-y", "-loglevel", "quiet"],
                timeout=60, check=True
            )
            return f"오디오 자르기 완료: {out.name}"
        except Exception as e:
            return f"자르기 실패 (ffmpeg 필요): {e}"

    elif action == "convert":
        fmt = params.get("format", "mp3")
        out = path.parent / f"{path.stem}.{fmt}"
        try:
            subprocess.run(
                ["ffmpeg", "-i", str(path), str(out), "-y", "-loglevel", "quiet"],
                timeout=60, check=True
            )
            return f"오디오 변환 완료: {out.name}"
        except Exception as e:
            return f"변환 실패 (ffmpeg 필요): {e}"

    return f"지원하지 않는 오디오 작업: {action}"


# ─── 비디오 처리 ─────────────────────────────────────────────────────────────

def _process_video(path: Path, action: str, params: dict) -> str:
    def _ffmpeg(*args, timeout=120):
        subprocess.run(
            ["ffmpeg"] + list(args) + ["-y", "-loglevel", "quiet"],
            timeout=timeout, check=True
        )

    if action == "trim":
        start = params.get("start", "0")
        end   = params.get("end", "")
        out   = path.parent / f"{path.stem}_trimmed{path.suffix}"
        end_args = ["-to", end] if end else []
        try:
            _ffmpeg("-i", str(path), "-ss", start, *end_args, "-c", "copy", str(out))
            return f"비디오 자르기 완료: {out.name}"
        except Exception as e:
            return f"자르기 실패 (ffmpeg 필요): {e}"

    elif action == "extract_audio":
        out = path.parent / f"{path.stem}.mp3"
        try:
            _ffmpeg("-i", str(path), "-vn", "-acodec", "mp3", str(out))
            return f"오디오 추출 완료: {out.name}"
        except Exception as e:
            return f"추출 실패 (ffmpeg 필요): {e}"

    elif action == "extract_frame":
        t   = params.get("time", "00:00:01")
        out = path.parent / f"{path.stem}_frame.jpg"
        try:
            _ffmpeg("-i", str(path), "-ss", t, "-frames:v", "1", str(out))
            return f"프레임 추출 완료: {out.name}"
        except Exception as e:
            return f"추출 실패 (ffmpeg 필요): {e}"

    elif action == "compress":
        crf = params.get("crf", "28")
        out = path.parent / f"{path.stem}_compressed.mp4"
        try:
            _ffmpeg("-i", str(path), "-vcodec", "libx264", "-crf", str(crf), str(out))
            orig = path.stat().st_size
            comp = out.stat().st_size
            saved = (1 - comp / orig) * 100 if orig else 0
            return f"비디오 압축 완료: {out.name} ({saved:.1f}% 절약)"
        except Exception as e:
            return f"압축 실패 (ffmpeg 필요): {e}"

    return f"지원하지 않는 비디오 작업: {action}"


# ─── 압축 파일 처리 ──────────────────────────────────────────────────────────

def _process_archive(path: Path, action: str, params: dict) -> str:
    import zipfile, tarfile

    if action == "list":
        try:
            if path.suffix.lower() == ".zip":
                with zipfile.ZipFile(path) as zf:
                    names = zf.namelist()
            elif path.suffix.lower() in (".tar", ".gz", ".bz2", ".xz"):
                with tarfile.open(path) as tf:
                    names = tf.getnames()
            else:
                return f"지원하지 않는 압축 형식: {path.suffix}"
            return f"파일 목록 ({len(names)}개):\n" + "\n".join(names[:50])
        except Exception as e:
            return f"목록 조회 실패: {e}"

    elif action == "extract":
        out_dir = path.parent / path.stem
        try:
            if path.suffix.lower() == ".zip":
                with zipfile.ZipFile(path) as zf:
                    zf.extractall(out_dir)
            elif path.suffix.lower() in (".tar", ".gz", ".bz2", ".xz"):
                with tarfile.open(path) as tf:
                    tf.extractall(out_dir)
            else:
                return f"지원하지 않는 압축 형식: {path.suffix}"
            return f"압축 해제 완료: {out_dir}"
        except Exception as e:
            return f"압축 해제 실패: {e}"

    return f"지원하지 않는 아카이브 작업: {action}"


# ─── PPTX 처리 ───────────────────────────────────────────────────────────────

def _process_pptx(path: Path, action: str, params: dict) -> str:
    _ensure("python-pptx")
    from pptx import Presentation

    prs   = Presentation(str(path))
    texts = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            texts.append(f"[슬라이드 {i+1}]\n" + "\n".join(slide_text))

    full_text = "\n\n".join(texts)

    if action == "summarize":
        return _call_claude(
            f"이 프레젠테이션을 한국어로 요약하세요:\n\n{full_text[:8000]}"
        ) or "요약 실패"
    elif action == "extract_text":
        return full_text or "텍스트를 찾을 수 없습니다."
    else:
        return _call_claude(
            f"이 프레젠테이션을 분석하세요:\n\n{full_text[:8000]}"
        ) or "분석 실패"


# ─── 메인 진입점 ─────────────────────────────────────────────────────────────

def file_processor(parameters: dict, player=None, speak=None) -> str:
    """
    파일 처리기 — 이미지, PDF, Word, Excel, 코드, 오디오, 비디오 등 처리

    parameters:
        path    : (필수) 처리할 파일 경로
        action  : (선택) 수행할 작업 (기본: 자동 감지)
        format  : (선택) 출력 포맷
        질문     : (선택) 파일에 대해 묻고 싶은 질문
    """
    params = parameters or {}
    path_str = params.get("path", "").strip()
    action   = params.get("action", "").strip().lower()
    question = params.get("question", params.get("질문", "")).strip()

    if not path_str:
        return "처리할 파일 경로를 지정해주세요."

    path = Path(path_str)
    if not path.exists():
        return f"파일을 찾을 수 없습니다: {path_str}"
    if not path.is_file():
        return f"파일이 아닙니다: {path_str}"

    cat = _cat(path)
    size_mb = path.stat().st_size / (1024 * 1024)

    if player:
        player.write_log(f"[파일] {path.name} ({cat}, {size_mb:.1f}MB)")

    print(f"[파일처리] {path.name} | 카테고리={cat} | 작업={action or '자동'}")

    # 특정 질문이 있으면 내용 읽고 Claude에게 질문
    if question and cat not in ("image", "audio", "video", "archive"):
        try:
            if cat == "pdf":
                content = _pdf_extract_text(path)
            elif cat in ("word", "text"):
                content = _read_text_file(path)
            elif cat in ("csv", "excel"):
                content = _read_csv(path)
            elif cat == "json":
                content = path.read_text(encoding="utf-8", errors="replace")
            elif cat == "code":
                content = path.read_text(encoding="utf-8", errors="replace")
            else:
                content = path.read_text(encoding="utf-8", errors="replace")
            return _call_claude(
                f"파일: {path.name}\n내용:\n{content[:8000]}\n\n질문: {question}"
            ) or "답변을 생성할 수 없습니다."
        except Exception as e:
            return f"파일 읽기 실패: {e}"

    # 카테고리별 처리
    try:
        if cat == "image":
            return _process_image(path, action or "describe", params)
        elif cat == "pdf":
            return _process_pdf(path, action or "summarize", params)
        elif cat in ("word", "text"):
            return _process_text(path, action or "summarize", params)
        elif cat in ("csv", "excel"):
            return _process_spreadsheet(path, action or "analyze", params)
        elif cat == "json":
            return _process_json(path, action or "validate", params)
        elif cat == "code":
            return _process_code(path, action or "explain", params)
        elif cat == "audio":
            return _process_audio(path, action or "transcribe", params)
        elif cat == "video":
            return _process_video(path, action or "extract_frame", params)
        elif cat == "archive":
            return _process_archive(path, action or "list", params)
        elif cat == "pptx":
            return _process_pptx(path, action or "summarize", params)
        else:
            # 알 수 없는 파일 — 텍스트로 읽기 시도
            try:
                content = path.read_text(encoding="utf-8", errors="replace")
                return _call_claude(
                    f"파일 '{path.name}'의 내용을 분석하세요:\n\n{content[:6000]}"
                ) or "분석할 수 없습니다."
            except Exception:
                return f"지원하지 않는 파일 형식: {path.suffix}"

    except Exception as e:
        print(f"[파일처리] ❌ {e}")
        return f"파일 처리 실패: {e}"
