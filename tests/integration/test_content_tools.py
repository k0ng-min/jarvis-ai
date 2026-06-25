"""YouTube routing and office document extraction tests."""

import tempfile
from pathlib import Path

import fitz
import openpyxl
from docx import Document
from pptx import Presentation

from jarvis_ai.actions.file_processor import file_processor
from jarvis_ai.actions.youtube_video import _extract_video_id
from jarvis_ai.main import _fast_route


def run():
    tool, params, _ = _fast_route(
        "유튜브 https://www.youtube.com/watch?v=dQw4w9WgXcQ 영상 분석해줘"
    )
    assert tool == "youtube_video"
    assert params["action"] == "analyze"
    assert _extract_video_id(params["url"]) == "dQw4w9WgXcQ"

    tool, params, _ = _fast_route("PDF 파일 읽어줘")
    assert tool == "file_processor"
    assert params["action"] == "extract_text"

    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)

        pdf_path = root / "sample.pdf"
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Jarvis PDF content")
        pdf.save(pdf_path)
        assert "Jarvis PDF content" in file_processor({
            "file_path": str(pdf_path),
            "action": "extract_text",
        })

        docx_path = root / "sample.docx"
        doc = Document()
        doc.add_paragraph("Jarvis Word content")
        doc.save(docx_path)
        assert "Jarvis Word content" in file_processor({
            "file_path": str(docx_path),
            "action": "extract_text",
        })

        xlsx_path = root / "sample.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["name", "value"])
        workbook.active.append(["Jarvis Excel", 42])
        workbook.save(xlsx_path)
        assert "Jarvis Excel" in file_processor({
            "file_path": str(xlsx_path),
            "action": "extract_text",
        })

        pptx_path = root / "sample.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Jarvis PowerPoint"
        slide.placeholders[1].text = "Slide content"
        deck.save(pptx_path)
        assert "Jarvis PowerPoint" in file_processor({
            "file_path": str(pptx_path),
            "action": "extract_text",
        })

    print("유튜브 분석 라우팅 및 PDF·Word·Excel·PowerPoint 읽기 통과")


if __name__ == "__main__":
    run()
